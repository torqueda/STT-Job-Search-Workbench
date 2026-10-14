#!/usr/bin/env python3
"""resume_generator.py: build editable resumes from approved canonical persona profiles.

Stable contact, education and experience come from base_resume.json. Projects, skills and coursework resolve
through canonical_resumes.json and evidence/claims.json (see workflow_quality.resolve_version). A batch spec
chooses which persona each company's resume uses and may reorder its projects; it cannot replace content.
Legacy specs and stale revisions fail on purpose.

Usage:
    python3 scripts/resume_generator.py --spec resume_versions/2026-01-05.json --dry-run
    python3 scripts/resume_generator.py --spec resume_versions/2026-01-05.json
    python3 scripts/resume_generator.py --spec resume_versions/2026-01-05.json --list

--dry-run validates and previews without writing and works before the spec is approved. Writing requires
"approved": true in the spec, which records that you made your Gate 1 choice. Output is one DOCX per persona
version. Open each file, check it meets your chosen page limit, and export it to PDF yourself (see docs/RESUME_HANDLING.md).
"""
import argparse
import json
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from workflow_quality import content_issues, require_v2, resolve_version  # noqa: E402
from workspace import announce, load_config_or_empty, read_json, workspace_dir  # noqa: E402

DEFAULT_LIMITS = {'max_skill_groups': 4, 'max_skills_total': 28, 'max_courses': 6, 'max_projects': 3,
                  'max_bullets_per_project': 2, 'max_bullets_per_role': 2}
DEFAULT_STYLE = {'font': 'Calibri', 'body_size': 10.5, 'heading_color': '0F243E',
                 'projects_heading': 'Projects', 'experience_heading': 'Professional Experience'}
UNSAFE_FILENAME_CHARS = set('/\\:*?"<>|')


def clean_filename_part(value, fallback):
    cleaned = ''.join(c for c in str(value).strip() if c not in UNSAFE_FILENAME_CHARS).strip('. ')
    return cleaned or fallback


def output_filename(base, kind, label):
    """'<Name> - <Kind> - <Label>.docx': one file per persona version, not per company."""
    return f"{clean_filename_part(base['name'], 'Name')} - {kind} - {clean_filename_part(label, 'Version')}.docx"


def _count_courses(line):
    if not line:
        return 0
    body = line.split(':', 1)[1] if ':' in line else line
    return len([c for c in body.split(',') if c.strip()])


def validate_experience(base, ws):
    """Stable experience bullets must meet the same evidence rule as persona projects."""
    ledger = read_json(Path(ws) / 'evidence' / 'claims.json')
    usable = {' '.join(c['text'].split()) for c in ledger.get('claims', {}).values()
              if c.get('status') == 'usable' and c.get('text')}
    for role in base.get('professional_experience', []):
        for bullet in role.get('bullets', []):
            if ' '.join(bullet.split()) not in usable:
                raise ValueError('Experience bullet is not a usable ledger claim: ' + bullet)


def validate_version(version, index, limits=None):
    """Human-readable warnings for one version. Never fatal: the user decides."""
    limits = {**DEFAULT_LIMITS, **(limits or {})}
    warns = []
    label = version.get('label') or f'#{index + 1}'
    skills = version.get('skills')
    if skills is not None:
        if not 2 <= len(skills) <= limits['max_skill_groups']:
            warns.append(f"[{label}] {len(skills)} skill groups (expected 2 to {limits['max_skill_groups']}).")
        total = sum(len(g.get('items', [])) for g in skills)
        if total > limits['max_skills_total']:
            warns.append(f"[{label}] {total} individual skills (limit {limits['max_skills_total']}).")
    projects = version.get('projects')
    if projects is not None:
        if not 2 <= len(projects) <= limits['max_projects']:
            warns.append(f"[{label}] {len(projects)} projects (expected 2 to {limits['max_projects']}).")
        for project in projects:
            count = len(project.get('bullets', []))
            if count > limits['max_bullets_per_project']:
                warns.append(f"[{label}] project '{project.get('title', '?')}' has {count} bullets "
                             f"(limit {limits['max_bullets_per_project']}).")
    coursework = version.get('coursework')
    if isinstance(coursework, str) and coursework.strip().lower() != 'omit':
        courses = _count_courses(coursework)
        if courses > limits['max_courses']:
            warns.append(f"[{label}] {courses} courses (limit {limits['max_courses']}).")
    return warns


def _docx_imports():
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
        from docx.shared import Inches, Pt, RGBColor
        return Document, Pt, Inches, RGBColor, WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
    except ImportError:
        sys.exit('python-docx is not installed. Run: python3 -m pip install python-docx')


def build_document(base, version, ws=None, style=None):
    validate_experience(base, workspace_dir(ws))
    Document, Pt, Inches, RGBColor, WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT = _docx_imports()
    style = {**DEFAULT_STYLE, **(base.get('style') or {}), **(style or {})}
    if 'canonical_id' in version and 'projects' not in version:
        version = resolve_version(version, ws=ws)
    errors, _ = content_issues(json.dumps(base) + json.dumps(version), ws)
    if errors:
        raise ValueError('\n'.join(errors))
    font, color = style['font'], style['heading_color']
    doc = Document()
    for section in doc.sections:
        section.top_margin = section.bottom_margin = Inches(0.5)
        section.left_margin = section.right_margin = Inches(0.6)
    right_tab = Inches(8.5 - 0.6 - 0.6)
    normal = doc.styles['Normal']
    normal.font.name = font
    normal.font.size = Pt(style['body_size'])
    normal.paragraph_format.space_after = normal.paragraph_format.space_before = Pt(0)

    def tight(paragraph, after=1, before=0):
        paragraph.paragraph_format.space_after = Pt(after)
        paragraph.paragraph_format.space_before = Pt(before)
        paragraph.paragraph_format.line_spacing = 1.0
        return paragraph

    def add_runs(paragraph, segments):  # segments: (text, bold, italic)
        for text, bold, italic in segments:
            run = paragraph.add_run(text)
            run.bold, run.italic, run.font.name = bold, italic, font
        return paragraph

    def titled_line(segments, date):
        paragraph = tight(doc.add_paragraph())
        paragraph.paragraph_format.tab_stops.add_tab_stop(right_tab, WD_TAB_ALIGNMENT.RIGHT)
        add_runs(paragraph, segments)
        if date:
            paragraph.add_run('\t' + date).font.name = font

    def heading(text):
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        paragraph = tight(doc.add_paragraph(), after=2, before=6)
        run = paragraph.add_run(text)
        run.bold, run.font.name, run.font.size = True, font, Pt(13)
        run.font.color.rgb = RGBColor.from_string(color)
        borders, bottom = OxmlElement('w:pBdr'), OxmlElement('w:bottom')
        for key, value in (('w:val', 'single'), ('w:sz', '6'), ('w:space', '1'), ('w:color', color)):
            bottom.set(qn(key), value)
        borders.append(bottom)
        paragraph._p.get_or_add_pPr().append(borders)

    def bullet(text):
        paragraph = tight(doc.add_paragraph())
        paragraph.paragraph_format.left_indent = Inches(0.2)
        paragraph.paragraph_format.first_line_indent = Inches(-0.15)
        paragraph.add_run('•  ' + text).font.name = font

    name = tight(doc.add_paragraph(), after=0)
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = name.add_run(base['name'])
    run.bold, run.font.name, run.font.size = True, font, Pt(20)
    run.font.color.rgb = RGBColor.from_string(color)
    contact = tight(doc.add_paragraph(), after=2)
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_run = contact.add_run(base['contact_line'])
    contact_run.font.name, contact_run.font.size = font, Pt(10)

    heading('Education')
    coursework = version.get('coursework', base.get('default_coursework'))
    for i, school in enumerate(base['education']):
        titled_line([(school['institution'], True, False)], school.get('date', ''))
        for key, italic in (('location', True), ('note', False)):
            if school.get(key):
                tight(add_runs(doc.add_paragraph(), [(school[key], False, italic)]))
        if i == 0 and coursework and str(coursework).strip().lower() != 'omit':  # coursework attaches to the newest degree
            tight(add_runs(doc.add_paragraph(), [(coursework, False, False)]))

    heading('Skills')
    for group in version.get('skills', base.get('default_skills', [])):
        add_runs(tight(doc.add_paragraph()), [(group['group'] + ': ', True, False), (', '.join(group.get('items', [])), False, False)])

    heading(style['projects_heading'])
    for project in version.get('projects', base.get('default_projects', [])):
        tight(add_runs(doc.add_paragraph(), [(project['title'], True, False)]), after=0)
        if project.get('purpose'):  # italic line: why the project existed, from a registered claim
            tight(add_runs(doc.add_paragraph(), [(project['purpose'], False, True)]), after=0)
        for text in project.get('bullets', []):
            bullet(text)

    heading(style['experience_heading'])
    for role in base.get('professional_experience', []):
        titled_line([(role['title'], True, False)], role.get('date', ''))
        for text in role.get('bullets', []):
            bullet(text)
    return doc


def main():
    parser = argparse.ArgumentParser(description='Generate editable .docx resumes per approved version.')
    parser.add_argument('--spec', required=True, help='Version spec JSON (schema_version 2).')
    parser.add_argument('--workspace')
    parser.add_argument('--base', help='Base resume JSON (default: <workspace>/base_resume.json).')
    parser.add_argument('--out-dir', help='Output folder (default: <workspace>/tailored_resumes).')
    parser.add_argument('--dry-run', action='store_true', help='Validate and preview; write nothing.')
    parser.add_argument('--list', action='store_true', help='List versions in the spec and exit.')
    args = parser.parse_args()

    ws = workspace_dir(args.workspace)
    cfg = load_config_or_empty(ws)
    announce(cfg)
    base = read_json(args.base or ws / 'base_resume.json')
    spec_path = Path(args.spec) if Path(args.spec).is_absolute() or Path(args.spec).exists() else ws / args.spec
    spec = read_json(spec_path)
    versions = spec.get('versions', [])
    if not versions:
        sys.exit("Spec contains no 'versions'.")
    if args.list:
        print(f"Spec: {spec_path.name}  |  approved={spec.get('approved', False)}  |  {len(versions)} version(s)")
        for i, v in enumerate(versions, 1):
            print(f"  {i}. {v.get('label', '?')} -> {', '.join(v.get('target_jobs', [])) or '(no jobs listed)'}")
        return
    try:
        require_v2(spec)
        validate_experience(base, ws)
        versions = [resolve_version(v, ws=ws) for v in versions]
    except ValueError as exc:
        sys.exit(str(exc))
    if not spec.get('approved', False) and not args.dry_run:
        sys.exit('GATE 1 not passed: the spec has "approved": false (or missing). No resume is written until you have '
                 'approved the triage buckets. Set "approved": true once you have. (--dry-run previews validation.)')
    warnings = []
    for i, version in enumerate(versions):
        warnings += validate_version(version, i, cfg.get('resume_limits'))
        if not version.get('companies'):
            warnings.append(f"[{version.get('label') or f'#{i + 1}'}] no \"companies\" list; add one so the batch records who it serves.")
    if warnings:
        print('Content-limit warnings:')
        for w in warnings:
            print('  ! ' + w)
        print()
    out_dir = Path(args.out_dir) if args.out_dir else ws / 'tailored_resumes'
    names = [output_filename(base, 'Resume', v['label']) for v in versions]
    if len(names) != len(set(names)):
        sys.exit('Duplicate output labels in this batch.')
    if args.dry_run:
        print(f'[DRY RUN] {len(versions)} file(s) would be generated into {out_dir}:')
        for version, name in zip(versions, names):
            print(f"  - {name}  ->  {', '.join(version.get('companies', []))}")
        return
    existing = [name for name in names if (out_dir / name).exists()]
    if existing:
        sys.exit('Refusing to overwrite existing resumes. Use a new label or output folder: ' + ', '.join(existing))
    out_dir.mkdir(parents=True, exist_ok=True)
    for version, name in zip(versions, names):
        build_document(base, version, ws=ws).save(out_dir / name)
        print(f"Generated {name}  ->  used for: {', '.join(version.get('companies', []))}")
    print(f'\nDone. {len(names)} resume file(s) in {out_dir}.')
    print('Next: review each file, export to PDF yourself, then submit (Gate 2).')


if __name__ == '__main__':
    main()
