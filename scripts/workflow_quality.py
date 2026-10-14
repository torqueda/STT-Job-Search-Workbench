"""Shared, offline content controls used by the generators, the answer check and package preparation.

Three kinds of control live here:

* Evidence resolution: resumes are assembled only from registered claims in the evidence ledger, through
  canonical persona profiles. A batch spec can select a profile and reorder its projects; it cannot replace
  its content (resolve_version).
* Known-error and editorial checks on text: a user-maintained list of statements that must never appear
  (evidence/known_errors.json) plus built-in voice warnings (content_issues, storytelling_issues).
* Review binding: a review is tied to the SHA-256 of whitespace-normalized extracted text, so wording edits invalidate
  it (validate_review).

These checks catch known mistakes and require a recorded review. They do not establish that a statement is
true or that writing sounds human; a person still reads the final text.
"""
import copy
import hashlib
import json
import re
from pathlib import Path

from workspace import load_config_or_empty, read_json, workspace_dir


def text_digest(text):
    return hashlib.sha256(' '.join(text.split()).encode()).hexdigest()


# --- Known errors (user-maintained) --------------------------------------------------------------

def load_known_errors(ws=None):
    path = workspace_dir(ws) / 'evidence' / 'known_errors.json'
    if not path.is_file():
        return []
    return read_json(path).get('rules', [])


def _known_error_hits(text, rules):
    """Each rule: {id, pattern, message, unless?}. `unless` is a regex that, when also present, clears the rule."""
    hits = []
    for rule in rules:
        if re.search(rule['pattern'], text, re.I) and not (rule.get('unless') and re.search(rule['unless'], text, re.I)):
            hits.append(rule['message'])
    return hits


# --- Editorial checks (built in, can be disabled in config.json) ---------------------------------

EDITORIAL = [
    ('gap_disclaimer', r"\bi (?:do not|don.t|have not|haven.t|lack|do lack)\b",
     'retain only if an actual requirement makes it useful; never a default closing.'),
    ('technical_dismissal', r"(?:easy|easier|straightforward) (?:part|half)|(?:model|technical work).*?(?:easy|straightforward)",
     'do not diminish the engineering contribution to set up a reflective lesson.'),
    ('contrast_formula', r"(?:not|wasn.t|isn.t)\b[^.!?]{0,150}(?:but|it was|it is)\b",
     'check for a forced not-X-but-Y construction; describe the concrete action instead.'),
    ('analogy', r"(?:same (?:habit|instinct|discipline)|at .*?scale|higher.stakes version)",
     'connect evidence to an actual responsibility; avoid a generic greater-scale conclusion.'),
    ('moral_comparison', r"peers|nobody was watching|no one was watching|no one was checking",
     'remove unsupported comparisons or moral lessons.'),
    ('adoption', r"real users|people started using|users started|a stranger would notice",
     'public deployment does not establish actual users; verify firsthand evidence.'),
    ('personal_reflection', r"\b(?:i|we)\b.*?(?:learned|realized|found|spent|enjoy)",
     'verify motivations, chronology, time spent and lessons with the user; do not infer them from project facts.'),
]


def content_issues(text, ws=None):
    """Return (errors, warnings). Errors block use; each warning needs a recorded resolution in a review."""
    errors = _known_error_hits(text, load_known_errors(ws))
    disabled = set(load_config_or_empty(ws).get('editorial_checks_disabled', []))
    lowered = text.casefold()
    warnings = [f'{key}: {message}' for key, pattern, message in EDITORIAL
                if key not in disabled and re.search(pattern, lowered)]
    return errors, warnings


# --- Storytelling check ---------------------------------------------------------------------------
# Applies to writing a stranger reads without the resume: screening answers, letter paragraphs, outreach
# and emails. Not run on resumes, whose project blocks carry registered purpose lines. Heuristics only:
# they catch a project named without saying what it was or why it existed, and stock motive phrasing.
# See prompts/09_storytelling_check.md.

KIND_CUE = r"\b(?:project|capstone|hackathon|course|class|assignment|study|thesis|client|internship|research)\b"
PURPOSE_CUE = (r"\b(?:wanted to|wants to|aim(?:ed|s)? (?:to|at)|goal|idea was|the point was|designed to|meant to|intended to|"
               r"built to|needed|so that|in order to|to help|to show|to test|to recommend|to route|to predict|to decide|"
               r"to find|to track|to identify|to understand|showing|shows|whether|question|trying to|tasked with|"
               r"looking to|test(?:ed|ing)? how)\b")
VAGUE_REFERENT = (r"\b(?:systems|projects|tools|work|products|software|solutions|applications|things|platforms) "
                  r"like (?:this|these|that|those)\b")
EMPLOYER_MISSION = (r"\b(?:for|at|with|within) (?:a|an|the) (?:\w+ ){0,2}(?:team|company|organization|organisation|business|"
                    r"startup|firm)s? (?:that|who|which) (?:is|are) (?:turning|transforming|revolutionizing|reimagining|"
                    r"redefining|democratizing|unlocking|changing)\b")
STOCK_PHRASE = (r"turning (?:\w+ ){0,3}(?:data|information|signals|insights|research) into "
                r"(?:decisions|action|value|insights|impact|outcomes)|at the intersection of|passionate about|"
                r"bridge the gap|cutting[- ]edge|fast[- ]paced|make an impact|make a difference")


def load_project_aliases(ws=None):
    path = workspace_dir(ws) / 'evidence' / 'project_aliases.json'
    if not path.is_file():
        return {}
    return {name: pattern for name, pattern in read_json(path).get('projects', {}).items()}


def _sentences(text):
    return [s for s in re.split(r'(?<=[.!?])\s+', ' '.join(text.split())) if s]


def storytelling_issues(text, ws=None):
    """Editorial warnings for answers and letters. Each key (text before the colon) needs a resolution."""
    warnings = []
    sentences = _sentences(text)
    lowered = [s.casefold() for s in sentences]
    for name, alias in load_project_aliases(ws).items():
        first = next((i for i, s in enumerate(lowered) if re.search(alias, s)), None)
        if first is None:
            continue
        window = ' '.join(lowered[max(0, first - 1):first + 3])
        has_kind = bool(re.search(KIND_CUE, window))
        has_purpose = bool(re.search(PURPOSE_CUE, window))
        if not (has_kind and has_purpose):
            missing = ' and '.join(label for label, ok in (
                ('what kind of project it was (personal, course, capstone, hackathon, client)', has_kind),
                ('why it existed or what it was for', has_purpose)) if not ok)
            warnings.append(
                f'project_intro: {name} is first mentioned ("{sentences[first][:70]}...") without {missing}. '
                'A reader who has never seen the resume must be able to say what it was and why it happened '
                'before the tools or numbers.')
    match = re.search(VAGUE_REFERENT, text, flags=re.I)
    if match:
        warnings.append(f'vague_referent: "{match.group(0)}" has no named antecedent; say what "this" is, '
                        'or restate the motive in your own words.')
    match = re.search(EMPLOYER_MISSION, text, flags=re.I)
    if match:
        warnings.append(f'employer_mission: "{match.group(0)}..." paraphrases the employer\'s mission in stock wording; '
                        'state your own motive (your growth, the responsibility you want to own) in first person.')
    match = re.search(STOCK_PHRASE, text, flags=re.I)
    if match:
        warnings.append(f'stock_phrase: "{match.group(0)}" is generic; replace it with the specific thing you did '
                        'or want to do.')
    return warnings


# --- Canonical resume resolution ------------------------------------------------------------------

BATCH_KEYS = {'canonical_id', 'canonical_revision', 'label', 'companies', 'target_jobs', 'job_family', 'project_order'}


def require_v2(spec):
    if spec.get('schema_version') != 2:
        raise ValueError('Legacy spec rejected. Use a schema_version=2 spec that references canonical_resumes.json. '
                         'Historical batches are not reusable authority.')


def _claim_text(ledger, claim_id, what):
    item = ledger['claims'].get(claim_id)
    if not item or item.get('status') != 'usable':
        raise ValueError(f'Unapproved {what}: {claim_id}')
    return item['text']


def resolve_version(version, catalog=None, ledger=None, ws=None):
    """Turn a batch version (profile reference plus optional ordering) into resume content from the ledger."""
    ws = workspace_dir(ws)
    catalog = catalog or read_json(ws / 'canonical_resumes.json')
    ledger = ledger or read_json(ws / 'evidence' / 'claims.json')
    extra = set(version) - BATCH_KEYS
    if extra:
        raise ValueError(f'Batch content overrides prohibited: {sorted(extra)}. Update the canonical evidence or profile '
                         'deliberately, with a change note, then reference its revision.')
    profile_id = version.get('canonical_id')
    if profile_id not in catalog['profiles']:
        raise ValueError(f'Unknown canonical resume: {profile_id}')
    if version.get('canonical_revision') != catalog['revision']:
        raise ValueError('Stale canonical revision. Re-read the current profile and evidence before preparing this batch.')
    profile = copy.deepcopy(catalog['profiles'][profile_id])
    if profile.get('status') == 'retired':
        raise ValueError(f'{profile_id} is retired ({profile.get("retired_note", "see change_note")}). '
                         'Use an active persona, or ask the user before reviving it.')
    if version.get('job_family', profile['family']) != profile['family']:
        raise ValueError('Job family does not match the canonical profile. Do not broaden a profile to absorb a different function.')
    projects_by_id = {p['id']: p for p in profile['projects']}
    order = version.get('project_order', [p['id'] for p in profile['projects']])
    if len(order) != len(set(order)) or set(order) != set(projects_by_id):
        raise ValueError('project_order must contain every canonical project once. Change project selections in the '
                         'canonical profile with an evidence-based change note.')
    used, projects = [], []
    for project_id in order:
        project = projects_by_id[project_id]
        bullets = []
        for ref in project['claim_ids']:
            bullets.append(_claim_text(ledger, ref, 'evidence claim'))
            used.append(ref)
        resolved = {'title': project['title'], 'bullets': bullets}
        if project.get('purpose_id'):  # italic line under the title: why the project existed; must be a registered claim
            resolved['purpose'] = _claim_text(ledger, project['purpose_id'], 'purpose claim')
            used.append(project['purpose_id'])
        projects.append(resolved)
    if not set(profile.get('required_claim_ids', [])).issubset(used):
        raise ValueError('Canonical profile lost required positioning evidence. Restore required claim IDs before generation.')
    profile.update(version)
    profile['label'] = version.get('label', version['canonical_id'])
    profile['theme'] = profile['core_positioning']
    if profile['label'] in catalog['profiles'] and profile['label'] != version['canonical_id']:
        raise ValueError('A variant cannot overwrite another canonical profile name.')  # filenames cannot be hijacked
    profile['projects'] = projects
    errors, _ = content_issues(json.dumps(profile), ws)
    if errors:
        raise ValueError('\n'.join(errors))
    return profile


# --- Review binding -------------------------------------------------------------------------------

def validate_review(text, review, require_company=False, writing=False, ws=None):
    """Return (errors, warnings) for a review object covering this normalized text."""
    ws = workspace_dir(ws)
    ledger = read_json(ws / 'evidence' / 'claims.json')
    errors, warnings = content_issues(text, ws)
    if writing:
        warnings = warnings + storytelling_issues(text, ws)
    if not isinstance(review, dict):
        return errors + ['Missing factual/editorial review. See prompts/08_review.md.'], warnings
    if review.get('text_sha256') != text_digest(text):
        errors.append('Review is stale: text changed after review. Review the actual current text and update its digest.')
    if review.get('evidence_revision') != ledger['revision']:
        errors.append('Review used an old evidence ledger.')
    if not review.get('reviewer') or not review.get('reviewed_at'):
        errors.append('Record reviewer and reviewed_at.')
    checks = review.get('claim_checks', [])
    if not checks:
        errors.append('Provide claim-by-claim source checks, including personal and company assertions.')
    known_ids = set()
    for section in ('claims', 'reflections', 'credentials', 'languages'):
        known_ids |= set(ledger.get(section, {}))
    for check in checks:
        if not (check.get('statement') and check.get('source') and check.get('scope_check')):
            errors.append('Each claim check needs statement, source and scope_check.')
        if check.get('claim_id') and check['claim_id'] not in known_ids:
            errors.append(f"Unknown claim ID: {check['claim_id']}")
    if require_company and not review.get('company_sources'):
        errors.append('Record verified employer/posting URLs and supported details in company_sources.')
    for field in ('role_match', 'voice', 'proofreading'):
        if not review.get(field):
            errors.append(f'Missing substantive {field} review notes.')
    resolutions = review.get('warning_resolutions', {})
    for warning in warnings:
        key = warning.split(':', 1)[0]
        if not resolutions.get(key):
            errors.append(f'Unresolved editorial check: {key}. Revise or record a specific reason for retaining it.')
    return errors, warnings


# --- Reading documents ----------------------------------------------------------------------------

def read_document(path):
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == '.docx':
        from docx import Document
        doc = Document(path)
        return '\n'.join(p.text for p in doc.paragraphs) + '\n' + '\n'.join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    if suffix == '.pdf':
        try:
            from pypdf import PdfReader
        except ImportError:
            raise SystemExit('Reading PDFs needs pypdf: python3 -m pip install pypdf')
        return '\n'.join(page.extract_text() or '' for page in PdfReader(path).pages)
    return path.read_text(encoding='utf-8')


if __name__ == '__main__':
    import argparse
    parser = argparse.ArgumentParser(description='Inspect final material and print its review digest. Never approves anything.')
    parser.add_argument('file')
    parser.add_argument('--workspace')
    parser.add_argument('--writing', action='store_true', help='Also run the storytelling check (answers, letters, emails).')
    args = parser.parse_args()
    ws = workspace_dir(args.workspace)
    body = read_document(args.file)
    problems, notes = content_issues(body, ws)
    if args.writing:
        notes = notes + storytelling_issues(body, ws)
    print(json.dumps({'text_sha256': text_digest(body), 'errors': problems, 'editorial_checks': notes}, indent=2))
    raise SystemExit(bool(problems))
