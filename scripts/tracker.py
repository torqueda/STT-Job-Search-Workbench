"""Tracker storage with two interchangeable backends: CSV (default, no dependencies) and Excel (.xlsx).

Every script talks to the tracker through the Tracker class, so the rest of the workflow does not care
which one the user chose. Both backends:

* expose rows as dicts keyed by the canonical FIELDS below, plus '_row' (the row number as the user would
  see it in a spreadsheet, header included);
* write atomically (temporary file, then replace);
* copy the file to <workspace>/backups/tracker/ before every write and keep only the newest `backup_keep`;
* refuse to write while another writer holds the lock file.

config.json selects the backend:

    "tracker": {"format": "csv",  "path": "tracker.csv"}
    "tracker": {"format": "xlsx", "path": "My Applications.xlsx", "sheet": "Sheet1", "header_row": 1,
                "columns": {"company": "B", "job_title": "C", ...}}     # columns optional: defaults to A..S
"""
import csv
import os
import shutil
import tempfile
from datetime import datetime
from pathlib import Path

from workspace import resolve

FIELDS = ['date', 'company', 'job_title', 'type', 'location', 'link', 'requisition_id', 'family', 'fit', 'source',
          'referral', 'resume_version', 'package', 'status', 'stage_date', 'response_date', 'interview_date',
          'next_action', 'notes']
HEADERS = {
    'date': 'Date applied', 'company': 'Company', 'job_title': 'Job title', 'type': 'Type', 'location': 'Location',
    'link': 'Link', 'requisition_id': 'Requisition ID', 'family': 'Job family', 'fit': 'Fit', 'source': 'Source',
    'referral': 'Referral', 'resume_version': 'Resume version', 'package': 'Application package', 'status': 'Status',
    'stage_date': 'Stage date', 'response_date': 'First response date', 'interview_date': 'Interview date',
    'next_action': 'Next action', 'notes': 'Notes',
}
DATE_FIELDS = {'date', 'stage_date', 'response_date', 'interview_date'}
HEADER_TO_FIELD = {v.casefold(): k for k, v in HEADERS.items()}


def parse_date(value):
    if value in (None, ''):
        return None
    if isinstance(value, datetime):
        return value
    text = str(value).strip()
    for fmt in ('%Y-%m-%d', '%Y-%m-%d %H:%M:%S', '%m/%d/%Y'):
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


def format_date(value):
    return value.strftime('%Y-%m-%d') if isinstance(value, datetime) else ('' if value is None else str(value))


class TrackerBusy(ValueError):
    pass


class Tracker:
    def __init__(self, cfg, ws):
        self.ws = Path(ws)
        spec = cfg.get('tracker', {'format': 'csv', 'path': 'tracker.csv'})
        self.format = spec.get('format', 'csv')
        self.path = resolve(self.ws, spec.get('path', 'tracker.csv'))
        self.sheet = spec.get('sheet')
        self.header_row = int(spec.get('header_row', 1))
        self.columns = spec.get('columns') or {f: chr(ord('A') + i) for i, f in enumerate(FIELDS)}
        self.backup_keep = int(cfg.get('backup_keep', 30))
        if self.format not in ('csv', 'xlsx'):
            raise ValueError(f"Unsupported tracker format: {self.format!r} (use 'csv' or 'xlsx').")

    # -- reading ---------------------------------------------------------------------------------

    def read(self):
        """All data rows that have a company or title, as dicts. Dates become datetimes."""
        if not self.path.exists():
            return []
        raw = self._read_csv() if self.format == 'csv' else self._read_xlsx()
        rows = []
        for number, values in raw:
            if not (values.get('company') or values.get('job_title')):
                continue
            row = {f: (parse_date(values.get(f)) if f in DATE_FIELDS else (values.get(f) or None)) for f in FIELDS}
            row['_row'] = number
            rows.append(row)
        return rows

    def _read_csv(self):
        with self.path.open(newline='', encoding='utf-8') as handle:
            lines = list(csv.reader(handle))
        if not lines:
            return []
        header = [HEADER_TO_FIELD.get(h.strip().casefold(), h.strip()) for h in lines[0]]
        return [(n, dict(zip(header, line))) for n, line in enumerate(lines[1:], start=2)]

    def _read_xlsx(self):
        import openpyxl
        from openpyxl.utils import column_index_from_string
        wb = openpyxl.load_workbook(self.path, read_only=True, data_only=True)
        ws = wb[self.sheet] if self.sheet else wb.active
        out = []
        for number, row in enumerate(ws.iter_rows(values_only=True), start=1):
            if number <= self.header_row:
                continue
            values = {}
            for field, letter in self.columns.items():
                index = column_index_from_string(letter) - 1
                values[field] = row[index] if len(row) > index else None
            out.append((number, values))
        wb.close()
        return out

    # -- writing ---------------------------------------------------------------------------------

    def apply(self, operations):
        """operations: [('append', {field: value}), ('update', row_number, {field: value})].

        Returns (row numbers touched, backup path or None). Raises TrackerBusy if another writer is active.
        """
        lock = self.path.with_suffix(self.path.suffix + '.writer-lock')
        if self.format == 'xlsx' and self.path.with_name('~$' + self.path.name).exists():
            raise TrackerBusy('Close the workbook in Excel before writing; its lock file is present.')
        self.path.parent.mkdir(parents=True, exist_ok=True)
        try:
            descriptor = os.open(lock, os.O_CREAT | os.O_EXCL | os.O_WRONLY)
        except FileExistsError:
            raise TrackerBusy('Another tracker writer is active; do not remove its lock while it runs.')
        try:
            backup = self._backup()
            touched = self._write_csv(operations) if self.format == 'csv' else self._write_xlsx(operations)
            return touched, backup
        finally:
            os.close(descriptor)
            lock.unlink(missing_ok=True)

    def _backup(self):
        if not self.path.exists():
            return None
        folder = self.ws / 'backups' / 'tracker'
        folder.mkdir(parents=True, exist_ok=True)
        target = folder / (self.path.name + datetime.now().strftime('.%Y%m%d-%H%M%S-%f.bak'))
        shutil.copy2(self.path, target)
        for old in sorted(folder.glob(self.path.name + '.*.bak'))[:-self.backup_keep]:
            old.unlink()
        return target

    def _write_csv(self, operations):
        lines = []
        if self.path.exists():
            with self.path.open(newline='', encoding='utf-8') as handle:
                lines = list(csv.reader(handle))
        header = [HEADERS[f] for f in FIELDS]
        if not lines:
            lines = [header]
        fields = [HEADER_TO_FIELD.get(h.strip().casefold(), h.strip()) for h in lines[0]]
        touched = []
        for op in operations:
            if op[0] == 'append':
                lines.append([self._csv_value(op[1].get(f)) for f in fields])
                touched.append(len(lines))
            else:
                _, number, changes = op
                for field, value in changes.items():
                    if field in fields:
                        lines[number - 1][fields.index(field)] = self._csv_value(value)
                touched.append(number)
        handle, temp = tempfile.mkstemp(suffix='.csv', dir=self.path.parent)
        try:
            with os.fdopen(handle, 'w', newline='', encoding='utf-8') as out:
                csv.writer(out).writerows(lines)
            os.replace(temp, self.path)
        finally:
            if os.path.exists(temp):
                os.unlink(temp)
        return touched

    @staticmethod
    def _csv_value(value):
        return format_date(value) if isinstance(value, datetime) else ('' if value is None else str(value))

    def _write_xlsx(self, operations):
        import openpyxl
        from openpyxl.utils import column_index_from_string
        if self.path.exists():
            wb = openpyxl.load_workbook(self.path)
            ws = wb[self.sheet] if self.sheet else wb.active
        else:
            wb = openpyxl.Workbook()
            ws = wb.active
            if self.sheet:
                ws.title = self.sheet
        for field, letter in self.columns.items():  # add missing headers, never overwrite a user's own
            cell = ws.cell(self.header_row, column_index_from_string(letter))
            if cell.value in (None, ''):
                cell.value = HEADERS[field]

        def put(row, field, value):
            if field not in self.columns:
                return
            cell = ws.cell(row, column_index_from_string(self.columns[field]))
            cell.value = value
            if isinstance(value, str):
                cell.data_type = 's'
            if isinstance(value, datetime):
                cell.number_format = 'yyyy-mm-dd'

        last = max([self.header_row] + [r for r in range(self.header_row + 1, ws.max_row + 1)
                                        if ws.cell(r, column_index_from_string(self.columns['company'])).value
                                        or ws.cell(r, column_index_from_string(self.columns['job_title'])).value])
        touched = []
        for op in operations:
            if op[0] == 'append':
                last += 1
                for field, value in op[1].items():
                    put(last, field, value)
                touched.append(last)
            else:
                _, number, changes = op
                for field, value in changes.items():
                    put(number, field, value)
                touched.append(number)
        handle, temp = tempfile.mkstemp(suffix='.xlsx', dir=self.path.parent)
        os.close(handle)
        try:
            wb.save(temp)
            wb.close()
            os.replace(temp, self.path)
        finally:
            if os.path.exists(temp):
                os.unlink(temp)
        return touched
