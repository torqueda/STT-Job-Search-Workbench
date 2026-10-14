"""Test suite for the job search workflow scripts.

Every test runs against a throw-away copy of the FICTIONAL demo workspace (examples/demo_persona/workspace),
so no test touches real data and none needs the network.

Run:  python3 -m unittest discover -s tests -v
"""
import argparse
import csv
import hashlib
import json
import shutil
import subprocess
import sys
import tempfile
import time
import unittest
from unittest.mock import patch
from urllib.error import HTTPError
from datetime import datetime, timedelta, timezone
from pathlib import Path

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / 'scripts'))

import answer_check  # noqa: E402
import archive_handoff  # noqa: E402
import check_company  # noqa: E402
import doctor  # noqa: E402
import employer_watch  # noqa: E402
import init_workspace  # noqa: E402
import pii_scan  # noqa: E402
import skip_list  # noqa: E402
import sourcing  # noqa: E402
import tracker_report  # noqa: E402
import tracker_update  # noqa: E402
from application_package import prepare, verify_package  # noqa: E402
from cover_letter_generator import validate_letter, letter_rules  # noqa: E402
from resume_generator import build_document, output_filename, validate_version  # noqa: E402
from tracker import Tracker, TrackerBusy  # noqa: E402
from workflow_quality import (content_issues, read_document, require_v2, resolve_version, storytelling_issues,  # noqa: E402
                              text_digest, validate_review)
from workspace import read_json, write_json, load_config  # noqa: E402

DEMO = REPO / 'examples' / 'demo_persona' / 'workspace'


class DemoWorkspaceCase(unittest.TestCase):
    """Gives each test its own mutable copy of the demo workspace."""

    def setUp(self):
        self._tmp = tempfile.TemporaryDirectory()
        self.ws = Path(self._tmp.name) / 'ws'
        shutil.copytree(DEMO, self.ws)
        self.catalog = read_json(self.ws / 'canonical_resumes.json')
        self.ledger = read_json(self.ws / 'evidence' / 'claims.json')
        self.cfg = load_config(self.ws)
        self.spec = {'canonical_id': 'DataAnalytics', 'canonical_revision': self.catalog['revision'],
                     'label': 'DataAnalytics', 'companies': ['Example'], 'target_jobs': ['Data Analyst - 123']}

    def tearDown(self):
        self._tmp.cleanup()

    def resolve(self, spec=None, catalog=None):
        return resolve_version(spec or self.spec, catalog, self.ledger, ws=self.ws)

    def review(self, text, **extra):
        review = {'text_sha256': text_digest(text), 'evidence_revision': self.ledger['revision'], 'reviewer': 'Test fixture',
                  'reviewed_at': '2026-01-06', 'claim_checks': [{'statement': text, 'source': 'ledger',
                  'scope_check': 'personal pipeline work', 'claim_id': 'bikeshare.how'}],
                  'company_sources': [{'url': 'https://example.com/job', 'detail': 'Test role'}],
                  'role_match': 'Forecast pipeline matches this fixture.', 'voice': 'Direct account; no generic closing.',
                  'proofreading': 'Checked fixture wording.', 'warning_resolutions': {}}
        review.update(extra)
        return review


class EvidenceResolutionTests(DemoWorkspaceCase):
    """Resumes come only from registered claims through canonical personas."""

    def test_legacy_spec_cannot_regenerate(self):
        with self.assertRaisesRegex(ValueError, 'Legacy'):
            require_v2({'approved': True, 'versions': []})

    def test_batch_cannot_replace_persona_content(self):
        self.spec['projects'] = [{'title': 'Invented', 'bullets': ['Built something.']}]
        with self.assertRaisesRegex(ValueError, 'overrides prohibited'):
            self.resolve()

    def test_reordering_keeps_required_claims(self):
        self.spec['project_order'] = ['pantry', 'bikeshare']
        result = self.resolve()
        self.assertEqual(result['projects'][0]['title'].split(' - ')[0], 'Pantry Planner')
        self.assertIn('41 to 29', result['projects'][1]['bullets'][1])

    def test_persona_that_loses_required_claim_is_blocked(self):
        self.catalog['profiles']['DataAnalytics']['projects'][0]['claim_ids'] = ['bikeshare.how']
        with self.assertRaisesRegex(ValueError, 'required positioning'):
            self.resolve(catalog=self.catalog)

    def test_stale_revision_and_family_mismatch(self):
        self.spec['canonical_revision'] = 'old'
        with self.assertRaisesRegex(ValueError, 'Stale'):
            self.resolve()
        self.spec['canonical_revision'] = self.catalog['revision']
        self.spec['job_family'] = 'ai_applications'
        with self.assertRaisesRegex(ValueError, 'family'):
            self.resolve()

    def test_label_cannot_overwrite_another_persona(self):
        self.spec['label'] = 'ProductAnalytics'
        with self.assertRaisesRegex(ValueError, 'overwrite'):
            self.resolve()

    def test_every_active_persona_resolves_cleanly(self):
        for name, profile in self.catalog['profiles'].items():
            if profile.get('status') == 'retired':
                continue
            result = self.resolve(dict(self.spec, canonical_id=name, label=name))
            self.assertTrue(result['projects'])
            self.assertFalse(content_issues(json.dumps(result), self.ws)[0])

    def test_retired_persona_is_refused(self):
        self.catalog['profiles']['DataAnalytics']['status'] = 'retired'
        with self.assertRaisesRegex(ValueError, 'retired'):
            self.resolve(catalog=self.catalog)

    def test_every_project_carries_a_registered_purpose_line(self):
        for name in self.catalog['profiles']:
            for project in self.resolve(dict(self.spec, canonical_id=name, label=name))['projects']:
                self.assertTrue(project.get('purpose'), f'{name}/{project["title"]} lacks a purpose line')

    def test_unregistered_purpose_is_blocked(self):
        self.catalog['profiles']['DataAnalytics']['projects'][0]['purpose_id'] = 'bikeshare.invented_purpose'
        with self.assertRaisesRegex(ValueError, 'purpose'):
            self.resolve(catalog=self.catalog)

    def test_claim_that_is_not_usable_is_blocked(self):
        self.ledger['claims']['bikeshare.result']['status'] = 'needs_confirmation'
        with self.assertRaisesRegex(ValueError, 'Unapproved evidence claim'):
            self.resolve()


class ContentCheckTests(DemoWorkspaceCase):
    """Known errors block; editorial checks demand a recorded resolution; reviews bind to exact text."""

    def test_known_error_rules_block_unsupported_claims(self):
        for text in ('Pantry Planner is used by thousands of users.', 'We reduced forecast error by 30% overall.'):
            with self.subTest(text=text):
                self.assertTrue(content_issues(text, self.ws)[0])

    def test_known_error_unless_clause_clears_a_qualified_statement(self):
        self.assertFalse(content_issues('Reduced error by 30% against the seasonal baseline.', self.ws)[0])
        self.assertTrue(content_issues('Reduced error by 30%.', self.ws)[0])

    def test_editorial_warnings_fire_and_can_be_disabled(self):
        text = "I do not have experience at your scale, but I work carefully when no one was watching."
        keys = {w.split(':')[0] for w in content_issues(text, self.ws)[1]}
        self.assertTrue({'gap_disclaimer', 'moral_comparison', 'analogy'} <= keys)
        self.cfg['editorial_checks_disabled'] = ['gap_disclaimer']
        write_json(self.ws / 'config.json', self.cfg)
        keys = {w.split(':')[0] for w in content_issues(text, self.ws)[1]}
        self.assertNotIn('gap_disclaimer', keys)

    def test_review_binds_to_the_actual_text(self):
        text = 'Built a feature pipeline for a rental forecast.'
        review = self.review(text)
        self.assertFalse(validate_review(text, review, True, ws=self.ws)[0])
        self.assertTrue(validate_review(text + ' Used by thousands.', review, True, ws=self.ws)[0])

    def test_editorial_warning_requires_a_recorded_resolution(self):
        text = "I do not have experience at your scale, but I work carefully."
        errors, warnings = validate_review(text, self.review(text), True, ws=self.ws)
        self.assertTrue(any('gap_disclaimer' in w for w in warnings))
        self.assertTrue(any('Unresolved editorial check: gap_disclaimer' in e for e in errors))
        resolved = self.review(text, warning_resolutions={'gap_disclaimer': 'The posting requires scale; stated plainly.',
                                                          'contrast_formula': 'Checked.'})
        self.assertFalse([e for e in validate_review(text, resolved, True, ws=self.ws)[0] if 'gap_disclaimer' in e])

    def test_review_may_cite_reflections_but_not_unknown_ids(self):
        text = 'Built the Pantry Planner because food kept going bad at home.'
        good = self.review(text)
        good['claim_checks'][0]['claim_id'] = 'pantry.motive'
        self.assertFalse(validate_review(text, good, True, ws=self.ws)[0])
        good['claim_checks'][0]['claim_id'] = 'made.up'
        self.assertTrue(any('Unknown claim ID' in e for e in validate_review(text, good, True, ws=self.ws)[0]))

    def test_stale_ledger_revision_invalidates_review(self):
        text = 'Built a pipeline.'
        review = self.review(text, evidence_revision='older')
        self.assertTrue(any('old evidence ledger' in e for e in validate_review(text, review, ws=self.ws)[0]))


class StorytellingTests(DemoWorkspaceCase):
    UNINTRODUCED = ('To talk about the model side, I evaluated a Library FAQ assistant with a FAISS index. I want to build '
                    'systems like this for a team that is turning information into decisions.')
    INTRODUCED = ('To talk about the model side, I evaluated a Library FAQ assistant in a course project. The question was '
                  'whether embedding retrieval could answer a university library\'s common questions better than keyword '
                  'search. I want to build systems like this as I grow, and own the work from first idea to hand-off.')

    def keys(self, text):
        return {w.split(':', 1)[0] for w in storytelling_issues(text, self.ws)}

    def test_project_without_purpose_or_kind_is_flagged(self):
        self.assertIn('project_intro', self.keys(self.UNINTRODUCED))
        self.assertIn('project_intro', self.keys('I built the Pantry Planner with Flask and SQLite.'))

    def test_introduced_project_passes_the_intro_check(self):
        self.assertNotIn('project_intro', self.keys(self.INTRODUCED))

    def test_stock_motive_phrases_are_flagged(self):
        self.assertTrue({'vague_referent', 'employer_mission', 'stock_phrase'} <= self.keys(
            'I am passionate about systems like this for a team that is turning data into decisions.'))

    def test_first_person_motive_is_not_flagged(self):
        self.assertFalse(self.keys(self.INTRODUCED) & {'employer_mission', 'stock_phrase'})

    def test_no_aliases_means_the_project_check_is_inactive(self):
        write_json(self.ws / 'evidence' / 'project_aliases.json', {'projects': {}})
        self.assertNotIn('project_intro', self.keys(self.UNINTRODUCED))

    def test_writing_review_requires_resolution_of_storytelling_flags_but_resumes_do_not(self):
        text = self.UNINTRODUCED
        review = self.review(text)
        errors, _ = validate_review(text, review, writing=True, ws=self.ws)
        self.assertTrue(any('Unresolved editorial check: project_intro' in e for e in errors))
        review['warning_resolutions'] = {k: 'reason recorded' for k in ('project_intro', 'vague_referent', 'employer_mission',
                                                                       'stock_phrase', 'personal_reflection')}
        self.assertEqual(validate_review(text, review, writing=True, ws=self.ws)[0], [])
        self.assertEqual(validate_review(text, dict(review, warning_resolutions={}), ws=self.ws)[0], [])

    def test_answer_extraction_prefers_blockquotes_then_long_paragraphs(self):
        quoted = '# Heading\n\n> First answer\n> continues here.\n\nNot an answer.\n\n> Second answer.\n'
        self.assertEqual(answer_check.extract_answers(quoted), ['First answer continues here.', 'Second answer.'])
        paragraphs = '# Heading\n\n- bullet\n\n' + ' '.join(['word'] * 45) + '\n\nshort line\n'
        self.assertEqual(len(answer_check.extract_answers(paragraphs)), 1)

    def test_demo_answers_one_clean_one_flagged(self):
        text = (self.ws / 'cover_letters' / '2026-01-06_demo' / 'screening_answers.md').read_text()
        report = answer_check.check_answers(answer_check.extract_answers(text), self.ws)
        self.assertEqual([answer_check.is_flagged(r) for r in report], [False, True])


class TrackerTests(DemoWorkspaceCase):
    def update(self, **kw):
        return tracker_update.update(tracker_update.make_args(**kw), self.cfg, self.ws)

    def new_app(self, **kw):
        base = dict(title='Data Analyst', company='Example Co', link='https://example.com/jobs/1', requisition_id='1',
                    date='2026-01-06', resume_spec='resume_versions/2026-01-06_demo.json', confirmed_submitted=True)
        base.update(kw)
        return base

    def test_url_normalization_keeps_requisition_but_drops_tracking(self):
        n = check_company.normalize_url
        self.assertEqual(n('https://x.com/job/123?utm_source=test&ref=board'), n('https://x.com/job/123'))
        self.assertNotEqual(n('https://x.com/jobs?id=123'), n('https://x.com/jobs?id=124'))

    def test_csv_round_trip_creates_a_backup_before_writing(self):
        tracker = Tracker(self.cfg, self.ws)
        before = len(tracker.read())
        touched, backup = tracker.apply([('append', {'company': 'Zed', 'job_title': 'Analyst', 'date': datetime(2026, 1, 7)})])
        self.assertEqual(len(tracker.read()), before + 1)
        self.assertEqual(tracker.read()[-1]['_row'], touched[0])
        self.assertTrue(backup.exists())

    def test_backups_are_pruned_to_the_configured_count(self):
        self.cfg['backup_keep'] = 3
        tracker = Tracker(self.cfg, self.ws)
        for n in range(6):
            tracker.apply([('append', {'company': f'Co{n}', 'job_title': 'Analyst'})])
        self.assertEqual(len(list((self.ws / 'backups' / 'tracker').glob('*.bak'))), 3)

    def test_second_writer_is_refused_while_the_lock_exists(self):
        tracker = Tracker(self.cfg, self.ws)
        (tracker.path.with_suffix('.csv.writer-lock')).write_text('')
        with self.assertRaises(TrackerBusy):
            tracker.apply([('append', {'company': 'Zed'})])

    def test_xlsx_backend_preserves_unrelated_cells_and_dates(self):
        import openpyxl
        path = self.ws / 'book.xlsx'
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = 'Sheet1'
        ws['A1'] = 'my own title cell'
        ws['Z5'] = 'unrelated note'
        wb.save(path)
        self.cfg['tracker'] = {'format': 'xlsx', 'path': 'book.xlsx', 'sheet': 'Sheet1', 'header_row': 2,
                               'columns': {'company': 'B', 'job_title': 'C', 'date': 'D', 'link': 'E', 'status': 'F'}}
        tracker = Tracker(self.cfg, self.ws)
        tracker.apply([('append', {'company': 'Zed', 'job_title': 'Analyst', 'date': datetime(2026, 1, 7), 'status': 'Submitted'})])
        tracker.apply([('update', 3, {'status': 'Interview'})])
        out = openpyxl.load_workbook(path).active
        self.assertEqual((out['A1'].value, out['Z5'].value), ('my own title cell', 'unrelated note'))
        self.assertEqual((out['B2'].value, out['B3'].value, out['F3'].value, out['D3'].value),
                         ('Company', 'Zed', 'Interview', datetime(2026, 1, 7)))
        self.assertEqual(tracker.read()[0]['status'], 'Interview')

    def test_new_application_requires_confirmation_and_an_evidence_trail(self):
        with self.assertRaisesRegex(ValueError, 'confirm'):
            self.update(**self.new_app(confirmed_submitted=False))
        with self.assertRaisesRegex(ValueError, 'require --package'):
            self.update(**self.new_app(resume_spec=None))

    def test_submission_date_cannot_be_assumed(self):
        before = (self.ws / 'tracker.csv').read_bytes()
        with self.assertRaisesRegex(ValueError, 'actual submission date'):
            self.update(**self.new_app(date=None))
        self.assertEqual((self.ws / 'tracker.csv').read_bytes(), before)
        self.assertFalse((self.ws / 'application_events.jsonl').exists())

    def test_exclusion_end_date_is_inclusive_and_respects_as_of(self):
        self.cfg['excluded_companies'] = [{'name': 'Example', 'until': '2026-01-05'}]
        tracker = Tracker(self.cfg, self.ws)
        self.assertFalse(check_company.check('Example', self.cfg, tracker, now=datetime(2026, 1, 5, 18))[0])
        self.assertTrue(check_company.check('Example', self.cfg, tracker, now=datetime(2026, 1, 6))[0])

    def test_resume_spec_must_name_the_company(self):
        with self.assertRaisesRegex(ValueError, 'No version'):
            self.update(**self.new_app(company='Nobody Inc.'))
        result = self.update(**self.new_app(company='Brightwater Example Inc.', link='https://careers.example.com/brightwater/jobs/1003',
                                            requisition_id='1003'))
        row = Tracker(self.cfg, self.ws).read()[-1]
        self.assertEqual((row['resume_version'], row['status']), ('DataAnalytics', 'Submitted'))
        self.assertTrue((self.ws / 'application_events.jsonl').exists())
        self.assertEqual(result['row'], row['_row'])

    def test_duplicate_url_and_requisition_are_fatal_but_the_company_cap_only_warns(self):
        args = self.new_app(company='Brightwater Example Inc.', link='https://careers.example.com/brightwater/jobs/1003?utm_x=1', requisition_id='1003')
        self.update(**args)
        with self.assertRaisesRegex(ValueError, 'URL already recorded'):
            self.update(**dict(args, link='https://careers.example.com/brightwater/jobs/1003'))
        with self.assertRaisesRegex(ValueError, 'requisition already recorded'):
            self.update(**dict(args, link='https://careers.example.com/brightwater/jobs/9999'))
        result = self.update(**dict(args, link='https://careers.example.com/brightwater/jobs/7777', requisition_id='7777'))
        self.assertTrue(result['cap_warnings'])  # recorded anyway: refusing would not un-send the application

    def test_incomplete_attempts_do_not_count_or_need_confirmation(self):
        self.update(title='Analyst', company='Zed', link='https://z.example.com/1', status='Incomplete - form timed out')
        ok, _ = check_company.check('Zed', self.cfg, Tracker(self.cfg, self.ws), max_per_company=1)
        self.assertTrue(ok)

    def test_stage_change_needs_a_date_but_rejection_defaults_to_today(self):
        with self.assertRaisesRegex(ValueError, 'stage-date'):
            self.update(update_row=2, status='Interview')
        today = datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
        self.update(update_row=2, status='Rejected')  # row 2 has no first-response date yet: it is filled with today
        self.update(update_row=3, status='Rejected')  # row 3 already has one (2025-12-02): it is kept
        rows = {r['_row']: r for r in Tracker(self.cfg, self.ws).read()}
        self.assertEqual((rows[2]['status'], rows[2]['stage_date'], rows[2]['response_date']), ('Rejected', today, today))
        self.assertEqual((rows[3]['stage_date'], rows[3]['response_date']), (today, datetime(2025, 12, 2)))

    def test_dry_run_changes_nothing(self):
        digest = hashlib.sha256((self.ws / 'tracker.csv').read_bytes()).hexdigest()
        result = self.update(update_row=2, status='Interview', stage_date='2026-01-10', dry_run=True)
        self.assertTrue(result['dry_run'])
        self.assertEqual(digest, hashlib.sha256((self.ws / 'tracker.csv').read_bytes()).hexdigest())
        self.assertFalse((self.ws / 'application_events.jsonl').exists())

    def test_update_keeps_identity_fields(self):
        self.update(update_row=3, status='Interview', stage_date='2026-01-10', interview_date='2026-01-12', next_action='Prep')
        row = next(r for r in Tracker(self.cfg, self.ws).read() if r['_row'] == 3)
        self.assertEqual((row['company'], row['job_title'], row['date'], row['next_action']),
                         ('Cedar Example Labs', 'Associate Data Scientist', datetime(2025, 11, 24), 'Prep'))

    def test_cohorts_never_call_a_blank_outcome_a_rejection(self):
        today = datetime(2026, 1, 15)
        rows = [{'date': today - timedelta(days=40), 'status': None}, {'date': today - timedelta(days=40), 'status': 'Offer'},
                {'date': today, 'status': 'Rejected'}]
        mature, groups = tracker_report.summarize(rows, today)
        counts = next(iter(groups.values()))
        self.assertEqual((len(mature), counts['no recorded outcome'], counts['recorded interview/offer outcomes']), (2, 1, 1))

    def test_company_cap_and_expiring_exclusion(self):
        tracker = Tracker(self.cfg, self.ws)
        self.assertFalse(check_company.check('Harbor Example Systems', self.cfg, tracker, now=datetime(2026, 1, 5))[0])
        self.assertTrue(check_company.check('Harbor Example Systems', self.cfg, tracker, now=datetime(2026, 6, 1))[0])
        self.assertFalse(check_company.check('Placeholder Robotics Example', self.cfg, tracker)[0])
        self.cfg['excluded_companies'] = [{'name': 'Old Co', 'until': '2020-01-01'}]
        self.assertTrue(check_company.check('Old Co', self.cfg, tracker)[0])


class SourcingTests(DemoWorkspaceCase):
    def listing(self, **kw):
        base = {'title': 'Junior Data Analyst', 'description': 'Collaborate with senior engineers with 10+ years of experience.',
                'company': 'Example', 'url': 'https://x/job/1', 'location': '', 'posted': None}
        base.update(kw)
        return base

    def test_failed_board_request_is_not_retried(self):
        for status in (429, 403, 503):
            with self.subTest(status=status):
                error = HTTPError('https://example.com/board', status, 'unavailable', {}, None)
                with patch('employer_watch.urllib.request.urlopen', side_effect=error) as request:
                    with self.assertRaises(HTTPError):
                        employer_watch.http_json('https://example.com/board')
                    self.assertEqual(request.call_count, 1)
                error.close()

    def test_sourcing_honors_expired_exclusions(self):
        self.cfg['excluded_companies'] = [{'name': 'Example', 'until': '2020-01-01'}]
        self.assertEqual(len(sourcing.apply_filters([self.listing()], self.cfg)[0]), 1)

    def test_watcher_skips_excluded_employers_without_a_request(self):
        self.cfg['excluded_companies'] = ['Example']
        write_json(self.ws / 'employer_watchlist.json', {'employers': [
            {'name': 'Example', 'ats': 'greenhouse', 'token': 'example'}]})
        def unexpected(_):
            self.fail('Excluded employer must not be fetched')
        candidates, statuses = employer_watch.watch(self.ws, self.cfg, 2, {'greenhouse': unexpected}, pause=0)
        self.assertEqual(candidates, [])
        self.assertEqual(statuses[0]['status'], 'excluded')

    def test_junior_posting_is_not_dropped_for_mentioning_senior_colleagues(self):
        kept, *_ = sourcing.apply_filters([self.listing()], self.cfg)
        self.assertEqual(len(kept), 1)
        kept, *_ = sourcing.apply_filters([self.listing(title='Senior Data Analyst')], self.cfg)
        self.assertFalse(kept)

    def test_reasons_for_dropping_are_counted_separately(self):
        listings = [self.listing(url='https://x/applied'), self.listing(url='https://x/skipped'),
                    self.listing(company='Placeholder Robotics Example', url='https://x/excl'),
                    self.listing(description='This role is unpaid', url='https://x/kw'), self.listing(url='https://x/ok')]
        kept, dropped, skipped, applied = sourcing.apply_filters(listings, self.cfg, {'https://x/skipped'}, {'https://x/applied?utm_a=1'})
        self.assertEqual((len(kept), dropped, skipped, applied), (1, 2, 1, 1))

    def test_lane_patterns_apply_to_feeds_only(self):
        off_lane = self.listing(title='Barista')
        self.assertEqual(len(sourcing.apply_filters([off_lane], self.cfg, use_lane=False)[0]), 1)
        self.assertEqual(len(sourcing.apply_filters([off_lane], self.cfg, use_lane=True)[0]), 0)

    def test_feed_rows_are_filtered_by_activity_and_age(self):
        now = time.time()
        rows = [{'company_name': 'A', 'title': 'Analyst', 'url': 'u1', 'locations': ['NYC', 'Remote'], 'date_posted': now - 3600, 'active': True, 'is_visible': True},
                {'company_name': 'B', 'title': 'Analyst', 'url': 'u2', 'date_posted': now - 3 * 86400, 'active': True},
                {'company_name': 'C', 'title': 'Analyst', 'url': 'u3', 'date_posted': now - 60, 'active': False}]
        self.assertEqual([l['company'] for l in sourcing.fetch_feed(rows, hours=24, now=now)], ['A'])
        self.assertEqual(len(sourcing.fetch_feed(rows, now=now)), 2)
        self.assertEqual(sourcing.fetch_feed(rows, hours=24, now=now)[0]['location'], 'NYC; Remote')

    def test_manual_blocks_parse_title_company_location_and_url(self):
        listings = sourcing.parse_manual('Data Analyst\nAcme\nBoston\nhttps://acme.example.com/1\nMore text\n===\nML Engineer\nBeta\n')
        self.assertEqual([(l['title'], l['company'], l['url']) for l in listings],
                         [('Data Analyst', 'Acme', 'https://acme.example.com/1'), ('ML Engineer', 'Beta', '')])

    def test_repeat_batches_preserve_previous_queries_and_content(self):
        first = sourcing.write_batch([self.listing()], 'data', 'query one', self.ws / 'b')
        before = first.read_bytes()
        again = sourcing.write_batch([self.listing(title='Different role')], 'data', 'query one', self.ws / 'b')
        other = sourcing.write_batch([self.listing()], 'data', 'query two', self.ws / 'b')
        self.assertTrue(again.name.endswith('_2.md'))
        self.assertTrue(other.name.endswith('_3.md'))
        self.assertEqual(first.read_bytes(), before)

    def test_batch_label_cannot_escape_output_directory(self):
        with self.assertRaisesRegex(ValueError, 'filename label'):
            sourcing.write_batch([self.listing()], '../outside', 'query', self.ws / 'b')
        self.assertFalse((self.ws / 'b').exists())

    def test_employer_watch_reads_greenhouse_and_lever_and_applies_filters(self):
        now = datetime(2026, 1, 10, tzinfo=timezone.utc)
        recent = '2026-01-09T10:00:00-05:00'
        gh = {'jobs': [
            {'title': 'Data Analyst', 'location': {'name': 'Springfield, ST'}, 'absolute_url': 'https://gh/1', 'first_published': recent, 'requisition_id': 'R1'},
            {'title': 'Senior Data Analyst', 'location': {'name': 'Springfield, ST'}, 'absolute_url': 'https://gh/2', 'first_published': recent},
            {'title': 'Data Analyst', 'location': {'name': 'Paris'}, 'absolute_url': 'https://gh/3', 'first_published': recent},
            {'title': 'Data Analyst', 'location': {'name': 'Springfield, ST'}, 'absolute_url': 'https://gh/4', 'first_published': '2025-11-01T00:00:00Z'},
            {'title': 'Warehouse driver', 'location': {'name': 'Springfield, ST'}, 'absolute_url': 'https://gh/5', 'first_published': recent}]}
        lever = [{'text': 'AI Engineer', 'categories': {'location': 'Remote - US'}, 'hostedUrl': 'https://lv/1', 'id': 'x',
                  'createdAt': int(datetime(2026, 1, 9, tzinfo=timezone.utc).timestamp() * 1000)}]
        write_json(self.ws / 'employer_watchlist.json', {'metros': {'home': r'\bSpringfield\b'}, 'employers': [
            {'name': 'GH Co', 'ats': 'greenhouse', 'token': 't', 'metro': 'home'},
            {'name': 'LV Co', 'ats': 'lever', 'token': 't', 'metro': 'home'},
            {'name': 'Manual Co', 'ats': 'manual', 'note': 'by hand'}, {'name': 'Broken Co', 'ats': 'greenhouse', 'token': 'bad'}]})

        def fetch(url, **_):
            if 'bad' in url:
                raise OSError('boom')
            return gh if 'greenhouse' in url else lever

        adapters = {'greenhouse': lambda e: employer_watch.adapter_greenhouse(e, fetch),
                    'lever': lambda e: employer_watch.adapter_lever(e, fetch)}
        candidates, statuses = employer_watch.watch(self.ws, self.cfg, 2, adapters, now=now, pause=0)
        self.assertEqual(sorted(c['url'] for c in candidates), ['https://gh/1', 'https://lv/1'])
        by_name = {s['employer']: s['status'] for s in statuses}
        self.assertEqual((by_name['Manual Co'], by_name['GH Co']), ('manual', 'ok'))
        self.assertTrue(by_name['Broken Co'].startswith('error'))  # one outage does not stop the run


class PackageAndDocumentTests(DemoWorkspaceCase):
    def make_spec(self, root, text='Built a feature pipeline for a rental forecast.', **review_extra):
        (root / 'resume.txt').write_text(text)
        (root / 'jd.txt').write_text('Example job description')
        review = self.review(text, **review_extra)
        (root / 'review.json').write_text(json.dumps(review))
        return {'company': 'Example', 'title': 'Data Analyst', 'requisition_id': '123', 'url': 'https://example.com/job',
                'family': 'data_analytics', 'fit': 'Good', 'source': 'employer', 'resume_version': 'DataAnalytics@2026-01-05.1',
                'date': '2026-01-06', 'jd_file': str(root / 'jd.txt'), 'fit_approved': True, 'fit_rationale': 'Direct forecast work',
                'responsibilities': [{'duty': 'Forecast', 'evidence': 'bikeshare.result'}] * 3,
                'attachments': [{'kind': 'resume', 'path': str(root / 'resume.txt'), 'review_file': str(root / 'review.json')}]}

    def test_package_preserves_files_and_detects_edits(self):
        root = Path(tempfile.mkdtemp(dir=self._tmp.name))
        spec = self.make_spec(root)
        path = prepare(spec, root / 'packages', self.ws)
        self.assertEqual(verify_package(path)['state'], 'prepared-not-submitted')
        with self.assertRaisesRegex(ValueError, 'already exists'):
            prepare(spec, root / 'packages', self.ws)
        (path / '01_resume.txt').write_text('Changed')
        with self.assertRaisesRegex(ValueError, 'changed'):
            verify_package(path)

    def test_package_resolves_workspace_relative_paths(self):
        root = self.ws / 'final'
        root.mkdir()
        spec = self.make_spec(root)
        spec['jd_file'] = 'final/jd.txt'
        spec['attachments'][0].update(path='final/resume.txt', review_file='final/review.json')
        package = prepare(spec, self.ws / 'application_packages', self.ws)
        self.assertEqual(verify_package(package)['files'][1]['archive_name'], '01_resume.txt')
        self.assertEqual(spec['jd_file'], 'final/jd.txt')  # caller's spec stays unchanged

    def test_resume_generation_preserves_existing_files(self):
        out = self.ws / 'tailored_resumes'
        out.mkdir(exist_ok=True)
        original = out / 'Jordan Example - Resume - DataAnalytics.docx'
        original.write_bytes(b'existing draft')
        run = subprocess.run([sys.executable, str(REPO / 'scripts' / 'resume_generator.py'),
                              '--workspace', str(self.ws), '--spec', 'resume_versions/2026-01-06_demo.json'],
                             capture_output=True, text=True)
        self.assertNotEqual(run.returncode, 0)
        self.assertIn('overwrite', run.stderr)
        self.assertEqual(original.read_bytes(), b'existing draft')
        self.assertFalse((out / 'Jordan Example - Resume - AIApplications.docx').exists())

    def test_unregistered_experience_bullet_blocks_generation(self):
        base = read_json(self.ws / 'base_resume.json')
        base['professional_experience'][0]['bullets'].append('Led ten analysts.')
        with self.assertRaisesRegex(ValueError, 'not a usable ledger claim'):
            build_document(base, self.spec, ws=self.ws)

    def test_package_honors_user_resume_page_limit(self):
        root = self.ws / 'final'
        root.mkdir()
        spec = self.make_spec(root)
        doc = build_document(read_json(self.ws / 'base_resume.json'), self.spec, ws=self.ws)
        path = root / 'resume.docx'
        doc.save(path)
        text = read_document(path)
        _, warnings = content_issues(text, self.ws)
        review = self.review(text, page_count=2, visual_review='Fictional test review; no layout assertion.',
                             warning_resolutions={w.split(':')[0]: 'Fictional test fixture.' for w in warnings})
        write_json(root / 'review.json', review)
        spec['attachments'][0]['path'] = str(path)
        with self.assertRaisesRegex(ValueError, 'page limit'):
            prepare(spec, self.ws / 'application_packages', self.ws)
        self.cfg['resume_limits']['max_pages'] = 2
        write_json(self.ws / 'config.json', self.cfg)
        self.assertTrue(prepare(spec, self.ws / 'application_packages', self.ws).is_dir())

    def test_package_refuses_without_gate_one_or_with_a_stale_review(self):
        root = Path(tempfile.mkdtemp(dir=self._tmp.name))
        spec = self.make_spec(root)
        with self.assertRaisesRegex(ValueError, 'Gate 1'):
            prepare(dict(spec, fit_approved=False), root / 'p', self.ws)
        (root / 'resume.txt').write_text('Edited after review.')
        with self.assertRaisesRegex(ValueError, 'stale'):
            prepare(spec, root / 'p', self.ws)

    def test_resume_document_contains_registered_purpose_and_respects_the_gate(self):
        base = read_json(self.ws / 'base_resume.json')
        doc = build_document(base, self.spec, ws=self.ws)
        text = '\n'.join(p.text for p in doc.paragraphs)
        self.assertIn('A course capstone for a city parks department', text)
        self.assertIn('Jordan Example', text)
        out = Path(self._tmp.name) / 'out'
        spec_path = self.ws / 'resume_versions' / '2026-01-06_demo.json'
        run = subprocess.run([sys.executable, str(REPO / 'scripts' / 'resume_generator.py'), '--workspace', str(self.ws),
                              '--spec', str(spec_path), '--out-dir', str(out)], capture_output=True, text=True)
        self.assertEqual(run.returncode, 0, run.stderr)
        self.assertEqual(sorted(p.name for p in out.glob('*.docx')),
                         ['Jordan Example - Resume - AIApplications.docx', 'Jordan Example - Resume - DataAnalytics.docx'])
        self.assertIn('Library FAQ Assistant', read_document(out / 'Jordan Example - Resume - AIApplications.docx'))
        spec = read_json(spec_path)
        spec['approved'] = False
        write_json(spec_path, spec)
        run = subprocess.run([sys.executable, str(REPO / 'scripts' / 'resume_generator.py'), '--workspace', str(self.ws),
                              '--spec', str(spec_path), '--out-dir', str(out / 'x')], capture_output=True, text=True)
        self.assertNotEqual(run.returncode, 0)
        self.assertIn('GATE 1', run.stderr)

    def test_legacy_spec_fails_in_the_generator(self):
        legacy = self.ws / 'resume_versions' / 'legacy.json'
        write_json(legacy, {'approved': True, 'versions': [{'label': 'x', 'projects': []}]})
        run = subprocess.run([sys.executable, str(REPO / 'scripts' / 'resume_generator.py'), '--workspace', str(self.ws),
                              '--spec', str(legacy), '--dry-run'], capture_output=True, text=True)
        self.assertNotEqual(run.returncode, 0)
        self.assertIn('Legacy', run.stderr)

    def test_resume_content_limits_warn_without_blocking(self):
        warns = validate_version({'label': 'X', 'projects': [{'title': 't', 'bullets': ['a', 'b', 'c']}] * 4}, 0)
        self.assertTrue(any('projects' in w for w in warns) and any('bullets' in w for w in warns))

    def test_letter_checks_flag_padding_phrases_and_configured_sentences(self):
        rules = letter_rules({'letters': {'required_closing': 'Thank you for your consideration.'}})
        warns = validate_letter({'company': 'Acme', 'body_paragraphs': ['I will hit the ground running.']}, 0, rules)
        text = ' '.join(warns)
        self.assertIn('banned phrase', text)
        self.assertIn('close with', text)
        self.assertIn('words', text)

    def test_output_filenames_strip_unsafe_characters(self):
        self.assertEqual(output_filename({'name': 'A B'}, 'Resume', 'Acme/Beta: ?'), 'A B - Resume - AcmeBeta.docx')


class WorkspaceAndHygieneTests(DemoWorkspaceCase):
    def test_init_workspace_copies_templates_and_never_overwrites(self):
        target = Path(self._tmp.name) / 'fresh'
        created, kept = init_workspace.init(target)
        self.assertIn('config.json', created)
        self.assertTrue((target / 'evidence' / 'claims.json').is_file())
        self.assertTrue((target / 'batches').is_dir())
        (target / 'PROFILE_AND_RULES.md').write_text('my edits')
        created, kept = init_workspace.init(target)
        self.assertEqual(created, [])
        self.assertEqual((target / 'PROFILE_AND_RULES.md').read_text(), 'my edits')

    def test_a_fresh_workspace_is_valid_but_warns_about_what_is_empty(self):
        target = Path(self._tmp.name) / 'fresh'
        init_workspace.init(target)
        errors, warnings = doctor.diagnose(target)
        self.assertEqual(errors, [])
        self.assertTrue(any('no claims yet' in w for w in warnings))

    def test_doctor_is_clean_on_the_demo_and_catches_broken_references(self):
        self.assertEqual(doctor.diagnose(self.ws), ([], []))
        self.catalog['profiles']['DataAnalytics']['projects'][0]['claim_ids'].append('no.such.claim')
        write_json(self.ws / 'canonical_resumes.json', self.catalog)
        errors, _ = doctor.diagnose(self.ws)
        self.assertTrue(any('unknown claim no.such.claim' in e for e in errors))

    def test_doctor_catches_invalid_regexes_and_unusable_claims(self):
        write_json(self.ws / 'evidence' / 'known_errors.json', {'rules': [{'id': 'bad', 'pattern': '(unclosed', 'message': 'm'}]})
        self.ledger['claims']['pantry.result']['status'] = 'needs_confirmation'
        write_json(self.ws / 'evidence' / 'claims.json', self.ledger)
        errors, _ = doctor.diagnose(self.ws)
        self.assertTrue(any('invalid regex' in e for e in errors))
        self.assertTrue(any('pantry.result, which is not usable' in e for e in errors))

    def test_doctor_warns_about_resume_lines_that_are_not_ledger_claims(self):
        base = read_json(self.ws / 'base_resume.json')
        base['professional_experience'][0]['bullets'].append('Led a team of ten analysts.')
        write_json(self.ws / 'base_resume.json', base)
        _, warnings = doctor.diagnose(self.ws)
        self.assertTrue(any('not a usable ledger claim' in w and 'Led a team' in w for w in warnings))

    def test_doctor_checks_filter_and_metro_regexes(self):
        self.cfg.update(excluded_title_patterns=['('], lane_title_patterns=['['], watch_filters={'lane': '('})
        write_json(self.ws / 'config.json', self.cfg)
        write_json(self.ws / 'employer_watchlist.json', {'metros': {'home': '('}, 'employers': []})
        errors, _ = doctor.diagnose(self.ws)
        for label in ('config.excluded_title_patterns', 'config.lane_title_patterns',
                      'config.watch_filters', 'employer_watchlist.metros'):
            self.assertTrue(any(label in error for error in errors), label)

    def test_pii_scan_checks_filenames_and_flags_unreadable_assets(self):
        root = Path(self._tmp.name) / 'release'
        root.mkdir()
        (root / ('person' + '@' + 'gmail.com.txt')).write_text('No contact in the body')
        (root / 'binary.dat').write_bytes(b'plain\x00bytes')
        (root / 'external.md').symlink_to(self.ws / 'PROFILE_AND_RULES.md')
        kinds = {item[2] for item in pii_scan.scan(root)}
        self.assertTrue({'email', 'binary file', 'symlink'} <= kinds)

    def test_skip_list_is_append_only_and_deduplicated(self):
        path = self.ws / 'skipped_listings.json'
        self.assertTrue(skip_list.add_entry(path, 'https://x/new', 'level stretch', company='Zed'))
        self.assertFalse(skip_list.add_entry(path, 'https://x/new', 'again'))
        self.assertIn('https://x/new', skip_list.load_skipped_urls(path))

    def test_archive_handoff_moves_only_older_entries(self):
        entries = ''.join(f'## 2026-01-{n:02d}\nDone: day {n}\n\n' for n in range(12, 0, -1))
        (self.ws / 'HANDOFF.md').write_text('# Handoff\n\nintro\n\n' + entries)
        moved, target = archive_handoff.archive(self.ws, keep=10)
        self.assertEqual(moved, 2)
        kept = (self.ws / 'HANDOFF.md').read_text()
        self.assertIn('intro', kept)
        self.assertIn('day 3', kept)
        self.assertNotIn('day 2', kept)
        self.assertIn('day 1', target.read_text())

    def test_pii_scan_flags_personal_data_and_allows_placeholders(self):
        root = Path(self._tmp.name) / 'repo'
        (root / 'workspace').mkdir(parents=True)
        # The fixtures below are assembled from pieces so this file itself passes the scan it tests.
        email, phone, ssn = 'real.person' + '@' + 'gmail.com', '212-555-' + '1234', '123-45-' + '6789'
        home, profile = '/Users/' + 'realname', 'linkedin.com/in/' + 'realhandle'
        (root / 'workspace' / 'private.md').write_text(email)  # the workspace is never scanned
        (root / 'ok.md').write_text('jordan.example@example.com, (000) 000-0000, you@example.com\n/Users/you/x')
        (root / 'bad.md').write_text(f'mail {email} or {phone}, SSN {ssn}, path {home}/docs, {profile}, key ghp_' + 'a' * 36)
        deny = Path(self._tmp.name) / 'deny.txt'
        deny.write_text('# comments are ignored\nSecretCorp\nre:Project\\s+Falcon\n')
        (root / 'notes.md').write_text('I interviewed at secretcorp and worked on project  falcon')
        found = pii_scan.scan(root, deny)
        kinds = {kind for relative, _, kind, _ in found if relative == 'bad.md'}
        self.assertTrue({'email', 'phone', 'ssn', 'home path', 'linkedin profile', 'api key'} <= kinds)
        self.assertEqual([f for f in found if f[0] == 'ok.md'], [])
        self.assertEqual(len([f for f in found if f[0] == 'notes.md']), 2)
        allowed = pii_scan.scan(root, deny, {'notes.md': {'SecretCorp'}})
        self.assertEqual(len([f for f in allowed if f[0] == 'notes.md']), 1)

    def test_repository_itself_passes_the_builtin_scan(self):
        self.assertEqual(pii_scan.scan(REPO), [])

    def test_every_demo_file_is_labelled_fictional(self):
        """A new user's assistant must never mistake the demo persona for real data."""
        self.assertTrue(read_json(DEMO / 'config.json').get('fictional_demo'))
        for path in sorted(DEMO.rglob('*')):
            if path.is_file() and path.name != '.DS_Store' and path.suffix in ('.json', '.md', '.txt', '.csv'):
                if path.name in ('tracker.csv', 'skipped_listings.json', 'application_events.jsonl'):
                    continue  # data files with no free text: labelled by the folder's README and config flag
                self.assertIn('FICTIONAL', path.read_text(), f'{path.relative_to(DEMO)} lacks a FICTIONAL marker')
        readme = (DEMO.parent / 'README.md').read_text()
        self.assertIn('FICTIONAL', readme)

    def test_scripts_announce_the_demo_workspace(self):
        run = subprocess.run([sys.executable, str(REPO / 'scripts' / 'doctor.py'), '--workspace', str(DEMO)],
                             capture_output=True, text=True)
        self.assertIn('FICTIONAL DEMO', run.stderr)


if __name__ == '__main__':
    unittest.main()
